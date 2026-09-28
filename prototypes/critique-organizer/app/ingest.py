"""Bounded parsers. Reject over-limit inputs rather than silently truncate."""
import hashlib
import io
import json
import mimetypes
import re
import zipfile
from datetime import datetime, timezone
from pathlib import PurePosixPath
from . import store

ALLOWED = {'.txt','.md','.csv','.json','.pdf','.docx','.xlsx'}
MAX_BYTES = 10 * 1024 * 1024
MAX_TEXT = 300000
MAX_BLOCKS = 5000


def safe_names(filename, relative_path):
    if not isinstance(filename,str) or not isinstance(relative_path,str):
        raise ValueError('파일명과 경로는 문자열이어야 합니다.')
    path = PurePosixPath(relative_path.replace('\\','/'))
    if path.is_absolute() or '..' in path.parts or ':' in relative_path or any(ord(c)<32 for c in filename+relative_path):
        raise ValueError('상대 경로만 허용합니다.')
    name = PurePosixPath(filename.replace('\\','/')).name
    if not name or len(name)>250 or len(relative_path)>1000:
        raise ValueError('파일명 또는 경로가 너무 깁니다.')
    return name,str(path)


def archive_guard(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries)>5000 or sum(e.file_size for e in entries)>40*1024*1024:
            raise ValueError('압축 해제 크기가 허용 범위를 초과합니다.')


def parse_file(filename, data):
    suffix = PurePosixPath(filename).suffix.lower()
    if suffix not in ALLOWED:
        raise ValueError('지원하지 않는 형식입니다: '+suffix)
    if len(data)>MAX_BYTES:
        raise ValueError('파일당 최대 10MB입니다. 파일 전체를 처리하지 않았습니다.')
    blocks=[]
    warnings=[]
    chars=0

    def append(text, locator):
        nonlocal chars
        if not text.strip():
            return
        chars+=len(text)
        if chars>MAX_TEXT or len(blocks)>=MAX_BLOCKS:
            raise ValueError('추출 텍스트 30만 자 또는 5,000개 블록 한도를 초과했습니다. 저장하지 않았습니다.')
        blocks.append({'text':text,'locator':locator})

    if suffix=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError('암호화된 PDF는 지원하지 않습니다.')
        if len(reader.pages)>100:
            raise ValueError('PDF는 최대 100페이지입니다. 저장하지 않았습니다.')
        empty=[]
        for page_number,page in enumerate(reader.pages,1):
            text=page.extract_text() or ''
            if not text.strip(): empty.append(page_number)
            append(text,{'kind':'page','page':page_number})
        if empty: warnings.append('텍스트가 없는 페이지(OCR 미지원): '+', '.join(map(str,empty)))
        warnings.append('PDF 텍스트 계층만 추출합니다. 이미지·도표 내용 및 추출 순서는 원본으로 확인하세요.')
        scope='pdf_text_layer'
    elif suffix=='.docx':
        archive_guard(data)
        from docx import Document
        doc=Document(io.BytesIO(data))
        for index,paragraph in enumerate(doc.paragraphs,1):
            append(paragraph.text,{'kind':'paragraph','paragraph':index})
        for ti,table in enumerate(doc.tables,1):
            for ri,row in enumerate(table.rows,1):
                for ci,cell in enumerate(row.cells,1):
                    append(cell.text,{'kind':'table_cell','table':ti,'row':ri,'column':ci})
        warnings.append('DOCX 본문 문단과 표 셀만 추출합니다. 머리말·꼬리말·텍스트 상자·주석·이미지는 제외됩니다.')
        scope='docx_body_paragraphs_and_tables'
    elif suffix=='.xlsx':
        archive_guard(data)
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(data),read_only=True,data_only=False)
        try:
            if len(book.worksheets)>10:
                raise ValueError('XLSX는 최대 10개 시트입니다. 저장하지 않았습니다.')
            for sheet in book.worksheets:
                if sheet.max_row>1000 or sheet.max_column>100:
                    raise ValueError('시트당 최대 1,000행·100열입니다. 저장하지 않았습니다.')
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            append(str(cell.value),{'kind':'cell','sheet':sheet.title,'cell':cell.coordinate})
        finally:
            book.close()
        warnings.append('XLSX 셀 값과 수식 문자열만 추출합니다. 계산·차트·이미지 해석은 수행하지 않습니다.')
        scope='xlsx_cell_values_and_formula_text'
    else:
        try:
            text=data.decode('utf-8-sig')
        except UnicodeDecodeError as exc:
            raise ValueError('UTF-8 텍스트 파일로 변환 후 올려 주세요. 깨진 텍스트는 저장하지 않습니다.') from exc
        for line_number,line in enumerate(text.splitlines(),1):
            append(line,{'kind':'line','line':line_number})
        scope='utf8_text_lines'
    if not blocks:
        raise ValueError('추출 가능한 텍스트가 없습니다. 스캔 문서는 OCR 후 올려 주세요.')
    return blocks, {'scope':scope,'block_count':len(blocks),'character_count':chars,'warnings':warnings,
                    'truncated':False,'original_preserved':True}


ORG_RE = re.compile(r'(?<![가-힣A-Za-z0-9])([가-힣A-Za-z0-9·]{2,30}(?:위원회|연구원|대학교|전자|주식회사|부|청))(?=(?:에서는|으로는|에서|에게|으로|은|는|이|가|을|를|의|와|과|도|에)?(?:\s|[.,!?·:;()]|$))')


def entity_candidates(text):
    # Surface names only. No person identity, semantic relationship or truth assertion.
    return sorted({m.group(1) for m in ORG_RE.finditer(text)})


def ingest(filename, relative_path, data):
    filename,relative_path=safe_names(filename,relative_path)
    blocks,info=parse_file(filename,data)
    doc_id=hashlib.sha256(data).hexdigest()
    mime=mimetypes.guess_type(filename)[0] or 'application/octet-stream'
    with store.connect() as db:
        db.execute('BEGIN IMMEDIATE')
        exists=db.execute('SELECT title,path,parse_info FROM documents WHERE id=?',(doc_id,)).fetchone()
        if exists:
            return {'id':doc_id,'title':exists['title'],'path':exists['path'],'duplicate':True,'parse_info':json.loads(exists['parse_info'])}
        db.execute('INSERT INTO documents VALUES(?,?,?,?,?,?,?)', (doc_id,filename,relative_path,mime,data,datetime.now(timezone.utc).isoformat(),json.dumps(info,ensure_ascii=False)))
        for i,block in enumerate(blocks):
            block_id=f'{doc_id}:{i}'
            db.execute('INSERT INTO blocks VALUES(?,?,?,?,?)',(block_id,doc_id,i,block['text'],json.dumps(block['locator'],ensure_ascii=False)))
            for name in entity_candidates(block['text']):
                entity_id='name:'+hashlib.sha256(name.encode()).hexdigest()
                db.execute('INSERT OR IGNORE INTO entities VALUES(?,?,?)',(entity_id,name,'organization_candidate'))
                db.execute('INSERT INTO mentions VALUES(?,?,?)',(doc_id,entity_id,block_id))
        db.execute("UPDATE metadata SET value=value+1 WHERE key='corpus_revision'")
    return {'id':doc_id,'title':filename,'path':relative_path,'duplicate':False,'parse_info':info}
